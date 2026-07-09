"""
Converts the markdown-ish text the LLM returns (headers, bullet lists,
pipe-delimited tables, bold text) into a downloadable .docx file.

This is deliberately a plain, dependency-light converter rather than a
general markdown engine -- it only needs to handle the shapes the
Literature Review and Paper Comparison prompts actually produce.
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt


def _add_inline_runs(paragraph, text: str) -> None:
    """Splits on **bold** markers and adds runs accordingly, so bold text
    from the LLM's markdown renders as actual bold in Word instead of
    showing literal asterisks."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def _is_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def _is_table_separator(line: str) -> bool:
    # e.g. "| --- | :---: | ---: |"
    return bool(re.match(r"^\s*\|[\s:|-]+\|\s*$", line))


def _parse_table_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def markdown_to_docx_bytes(title: str, markdown_text: str, subtitle: str | None = None) -> bytes:
    doc = Document()

    doc.add_heading(title, level=0)
    if subtitle:
        subtitle_para = doc.add_paragraph()
        subtitle_run = subtitle_para.add_run(subtitle)
        subtitle_run.italic = True
        subtitle_run.font.size = Pt(11)

    lines = markdown_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # --- Table block: consume consecutive pipe-delimited lines ---
        if _is_table_row(stripped):
            table_lines = []
            while i < len(lines) and _is_table_row(lines[i].strip()):
                table_lines.append(lines[i].strip())
                i += 1
            table_lines = [ln for ln in table_lines if not _is_table_separator(ln)]
            if table_lines:
                rows = [_parse_table_row(ln) for ln in table_lines]
                num_cols = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=num_cols)
                table.style = "Light Grid Accent 1"
                for r_idx, row_cells in enumerate(rows):
                    row = table.add_row()
                    for c_idx in range(num_cols):
                        text = row_cells[c_idx] if c_idx < len(row_cells) else ""
                        cell_para = row.cells[c_idx].paragraphs[0]
                        _add_inline_runs(cell_para, text)
                        if r_idx == 0:
                            for run in cell_para.runs:
                                run.bold = True
                doc.add_paragraph()
            continue

        # --- Headings ---
        heading_match = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if heading_match:
            level = min(len(heading_match.group(1)) + 1, 4)  # keep below the title's level 0
            doc.add_heading(heading_match.group(2), level=level)
            i += 1
            continue

        # --- Bullet list items ---
        bullet_match = re.match(r"^[-*]\s+(.*)", stripped)
        if bullet_match:
            para = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(para, bullet_match.group(1))
            i += 1
            continue

        # --- Numbered list items ---
        numbered_match = re.match(r"^\d+\.\s+(.*)", stripped)
        if numbered_match:
            para = doc.add_paragraph(style="List Number")
            _add_inline_runs(para, numbered_match.group(1))
            i += 1
            continue

        # --- Plain paragraph ---
        para = doc.add_paragraph()
        _add_inline_runs(para, stripped)
        i += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
